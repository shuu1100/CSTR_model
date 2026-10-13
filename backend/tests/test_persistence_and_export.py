"""Test suite for Phase 12: Persistence, History, Run Comparison, and Exports.

Adheres strictly to GLOBAL-INSTRUCTIONS.md, phase-12, and phase-08A:
- Verifies SQLite + SQLAlchemy 2.0 CRUD persistence across restarts.
- Asserts metadata round-trip integrity (parameters, summaries, time series, provenance).
- Verifies side-by-side run comparisons.
- Verifies standardized CSV and JSON exports with units, timestamps, and schema version.
"""

import json
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from backend.app.db.database import Base, get_db
from backend.app.db.service import (
    compare_runs,
    delete_simulation_run,
    export_runs_to_csv,
    export_runs_to_json,
    get_simulation_run,
    list_simulation_runs,
    record_simulation_run,
)
from backend.app.main import app

client = TestClient(app)


@pytest.fixture
def in_memory_db() -> Session:
    """Provide clean in-memory SQLite database session for unit tests."""
    engine = create_engine("sqlite:///:memory:", echo=False)
    Base.metadata.create_all(bind=engine)
    SessionTest = sessionmaker(bind=engine)
    session = SessionTest()
    try:
        yield session
    finally:
        session.close()


class TestDatabaseCRUD:
    """Unit tests for SQLite database operations."""

    def test_record_and_get_simulation_run(self, in_memory_db: Session) -> None:
        rec = record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            name="Canonical Oracle Test",
            parameters={"V": 10.0, "Q": 1.0, "k": 0.15, "C_in": 2.0, "C_0": 0.0},
            results_summary={"tau": 10.0, "C_ss": 0.8, "conversion": 0.6},
            time_series={"t": [0.0, 1.0, 2.0], "C_A": [0.0, 0.177, 0.315]},
            provenance="SIMULATED",
            metadata={"oracle": True, "units": {"V": "L", "Q": "L/min"}},
        )

        assert rec.id.startswith("run-")
        fetched = get_simulation_run(in_memory_db, rec.id)
        assert fetched is not None
        assert fetched.name == "Canonical Oracle Test"
        assert fetched.get_parameters()["V"] == 10.0
        assert fetched.get_results_summary()["C_ss"] == 0.8
        assert fetched.get_time_series()["t"] == [0.0, 1.0, 2.0]
        assert fetched.provenance == "SIMULATED"

    def test_list_and_filter_runs(self, in_memory_db: Session) -> None:
        record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            parameters={"V": 10.0},
            results_summary={"tau": 10.0},
        )
        record_simulation_run(
            db=in_memory_db,
            run_type="steady_state",
            parameters={"V": 20.0},
            results_summary={"tau": 20.0},
        )

        all_runs = list_simulation_runs(in_memory_db)
        assert len(all_runs) >= 2

        dynamic_runs = list_simulation_runs(in_memory_db, run_type="dynamic")
        assert all(r.run_type == "dynamic" for r in dynamic_runs)

    def test_delete_run(self, in_memory_db: Session) -> None:
        rec = record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            parameters={"V": 5.0},
            results_summary={"tau": 5.0},
        )
        assert delete_simulation_run(in_memory_db, rec.id) is True
        assert get_simulation_run(in_memory_db, rec.id) is None
        assert delete_simulation_run(in_memory_db, "non-existent-id") is False


class TestRunComparisonAndExports:
    """Tests for multi-run comparison and CSV/JSON export generation."""

    def test_compare_two_runs(self, in_memory_db: Session) -> None:
        r1 = record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            name="Run 1 (k=0.15)",
            parameters={"k": 0.15, "V": 10.0},
            results_summary={"C_ss": 0.8},
        )
        r2 = record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            name="Run 2 (k=0.30)",
            parameters={"k": 0.30, "V": 10.0},
            results_summary={"C_ss": 0.5},
        )

        res = compare_runs(in_memory_db, [r1.id, r2.id])
        assert res["compared_runs_count"] == 2
        assert len(res["runs"]) == 2
        assert res["runs"][0]["parameters"]["k"] in (0.15, 0.30)

    def test_compare_single_run_raises(self, in_memory_db: Session) -> None:
        r1 = record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            parameters={"k": 0.15},
            results_summary={},
        )
        with pytest.raises(ValueError, match="At least 2 valid run IDs are required"):
            compare_runs(in_memory_db, [r1.id])

    def test_export_json_valid_schema(self, in_memory_db: Session) -> None:
        record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            name="Export Test Run",
            parameters={"V": 10.0, "Q": 1.0},
            results_summary={"tau": 10.0},
            provenance="SIMULATED",
        )

        json_str = export_runs_to_json(in_memory_db)
        data = json.loads(json_str)
        assert data["schema_version"] == "1.0"
        assert "export_timestamp" in data
        assert len(data["runs"]) >= 1
        assert data["runs"][0]["name"] == "Export Test Run"

    def test_export_csv_valid_columns(self, in_memory_db: Session) -> None:
        record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            name="CSV Run",
            parameters={"V": 15.0, "Q": 1.5},
            results_summary={"tau": 10.0, "C_ss": 0.4},
            provenance="SIMULATED",
        )

        csv_str = export_runs_to_csv(in_memory_db)
        lines = csv_str.strip().split("\n")
        header = lines[0].split(",")
        assert "run_id" in header
        assert "run_type" in header
        assert "provenance" in header
        assert "param_V" in header
        assert "result_tau" in header


class TestHistoryAPIEndpoints:
    """Integration tests for FastAPI /api/v1/history routes."""

    def test_api_create_get_delete_cycle(self) -> None:
        # Create
        payload = {
            "run_type": "dynamic",
            "name": "API Integration Run",
            "parameters": {"V": 12.0, "Q": 1.2},
            "results_summary": {"tau": 10.0},
            "provenance": "SIMULATED",
        }
        create_res = client.post("/api/v1/history/runs", json=payload)
        assert create_res.status_code == 201
        run_id = create_res.json()["id"]

        # List
        list_res = client.get("/api/v1/history/runs")
        assert list_res.status_code == 200
        run_ids = [r["id"] for r in list_res.json()]
        assert run_id in run_ids

        # Get by ID
        get_res = client.get(f"/api/v1/history/runs/{run_id}")
        assert get_res.status_code == 200
        assert get_res.json()["name"] == "API Integration Run"

        # Delete
        del_res = client.delete(f"/api/v1/history/runs/{run_id}")
        assert del_res.status_code == 204

        # Confirm 404
        assert client.get(f"/api/v1/history/runs/{run_id}").status_code == 404

    def test_api_compare_runs(self) -> None:
        r1 = client.post("/api/v1/history/runs", json={
            "run_type": "dynamic", "parameters": {"k": 0.1}, "results_summary": {"C_ss": 0.8}
        }).json()["id"]
        r2 = client.post("/api/v1/history/runs", json={
            "run_type": "dynamic", "parameters": {"k": 0.2}, "results_summary": {"C_ss": 0.6}
        }).json()["id"]

        comp_res = client.post("/api/v1/history/compare", json={"run_ids": [r1, r2]})
        assert comp_res.status_code == 200
        data = comp_res.json()
        assert data["compared_runs_count"] == 2

    def test_api_export_csv_and_json(self) -> None:
        res_csv = client.get("/api/v1/history/export?format=csv")
        assert res_csv.status_code == 200
        assert "text/csv" in res_csv.headers["content-type"]
        assert "run_id" in res_csv.text

        res_json = client.get("/api/v1/history/export?format=json")
        assert res_json.status_code == 200
        assert "application/json" in res_json.headers["content-type"]
        data = res_json.json()
        assert "schema_version" in data

    def test_api_export_xlsx(self) -> None:
        import io
        import openpyxl

        res_xlsx = client.get("/api/v1/history/export?format=xlsx")
        assert res_xlsx.status_code == 200
        assert "openxmlformats-officedocument" in res_xlsx.headers["content-type"]

        wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
        assert "Runs_Overview" in wb.sheetnames
        assert "Parameters" in wb.sheetnames
        assert "Time_Series" in wb.sheetnames

    def test_api_retention_endpoint(self) -> None:
        res = client.post("/api/v1/history/retention/apply?max_runs=100")
        assert res.status_code == 200
        data = res.json()
        assert data["status"] == "success"
        assert "purged_count" in data


class TestPersistencePhase12Advanced:
    """Advanced unit tests for SQLite disk persistence, multi-sheet XLSX, and retention."""

    def test_sqlite_disk_persistence_across_restart(self, tmp_path) -> None:
        db_file = tmp_path / "cstr_test_disk.db"
        db_url = f"sqlite:///{db_file}"

        # 1. Initialize and write
        engine1 = create_engine(db_url)
        Base.metadata.create_all(bind=engine1)
        SessionMaker1 = sessionmaker(bind=engine1)
        session1 = SessionMaker1()

        rec = record_simulation_run(
            db=session1,
            run_type="dynamic",
            name="Disk Persistence Test",
            parameters={"V": 10.0, "Q": 1.0, "k": 0.15, "C_in": 2.0, "C_0": 0.0},
            results_summary={"tau": 10.0, "C_ss": 0.8},
            time_series={"t": [0.0, 1.0, 2.0], "C_A": [0.0, 0.177, 0.315]},
            provenance="SIMULATED",
            schema_version="1.0",
            artifact_ref="models/pinn_surrogate.pt",
            seed=42,
            split_id="split_80_20_seed42",
            metadata={"units": "canonical"},
        )
        saved_id = rec.id
        session1.close()
        engine1.dispose()

        # 2. Reconnect with a new engine to verify disk persistence across shutdown/restart
        engine2 = create_engine(db_url)
        SessionMaker2 = sessionmaker(bind=engine2)
        session2 = SessionMaker2()

        loaded = get_simulation_run(session2, saved_id)
        assert loaded is not None
        assert loaded.name == "Disk Persistence Test"
        assert loaded.schema_version == "1.0"
        assert loaded.artifact_ref == "models/pinn_surrogate.pt"
        assert loaded.seed == 42
        assert loaded.split_id == "split_80_20_seed42"
        assert loaded.get_parameters()["V"] == 10.0
        assert loaded.get_results_summary()["C_ss"] == 0.8
        assert loaded.get_time_series()["t"] == [0.0, 1.0, 2.0]
        assert loaded.get_time_series()["C_A"] == [0.0, 0.177, 0.315]

        session2.close()
        engine2.dispose()

    def test_export_runs_to_xlsx_multisheet(self, in_memory_db: Session) -> None:
        import io
        import openpyxl
        from backend.app.db.service import export_runs_to_xlsx

        record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            name="XLSX Export Test Run",
            parameters={"V": 10.0, "Q": 1.0, "k": 0.15},
            results_summary={"tau": 10.0, "C_ss": 0.8},
            time_series={"t": [0.0, 0.5, 1.0], "C_A": [0.0, 0.1, 0.177]},
            schema_version="1.0",
            seed=123,
            split_id="split_test",
        )

        xlsx_bytes = export_runs_to_xlsx(in_memory_db)
        assert len(xlsx_bytes) > 0

        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        assert set(wb.sheetnames) == {"Runs_Overview", "Parameters", "Time_Series"}

        overview_sheet = wb["Runs_Overview"]
        headers = [cell.value for cell in overview_sheet[1]]
        assert "Run ID" in headers
        assert "Name" in headers
        assert "Seed" in headers
        assert "Split ID" in headers

        params_sheet = wb["Parameters"]
        p_headers = [cell.value for cell in params_sheet[1]]
        assert "Parameter" in p_headers
        assert "Value" in p_headers

        ts_sheet = wb["Time_Series"]
        ts_headers = [cell.value for cell in ts_sheet[1]]
        assert "Time t [min]" in ts_headers
        assert "Concentration C_A [mol/L]" in ts_headers
        assert ts_sheet.max_row >= 4  # Header + 3 time points

    def test_export_filtering_by_type_and_date(self, in_memory_db: Session) -> None:
        from datetime import datetime, timezone
        from backend.app.db.service import export_runs_to_json

        r1 = record_simulation_run(
            db=in_memory_db,
            run_type="dynamic",
            name="Dynamic Run",
            parameters={"k": 0.1},
            results_summary={},
        )
        r2 = record_simulation_run(
            db=in_memory_db,
            run_type="steady_state",
            name="Steady State Run",
            parameters={"k": 0.2},
            results_summary={},
        )

        json_dynamic = json.loads(export_runs_to_json(in_memory_db, run_type="dynamic"))
        assert json_dynamic["total_runs"] == 1
        assert json_dynamic["runs"][0]["id"] == r1.id

        json_ss = json.loads(export_runs_to_json(in_memory_db, run_type="steady_state"))
        assert json_ss["total_runs"] == 1
        assert json_ss["runs"][0]["id"] == r2.id

    def test_apply_retention_policy_limits(self, in_memory_db: Session) -> None:
        from backend.app.db.service import apply_retention_policy, list_simulation_runs

        for i in range(5):
            record_simulation_run(
                db=in_memory_db,
                run_type="dynamic",
                name=f"Run {i}",
                parameters={"k": 0.1 * i},
                results_summary={},
            )

        assert len(list_simulation_runs(in_memory_db)) == 5

        purged = apply_retention_policy(in_memory_db, max_runs=3)
        assert purged == 2

        remaining = list_simulation_runs(in_memory_db)
        assert len(remaining) == 3

