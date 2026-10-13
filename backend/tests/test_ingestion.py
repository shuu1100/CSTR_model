"""Tests for CSV/Excel data ingestion, inspection, Savitzky-Golay smoothing, anomaly detection, and API endpoints."""

import io
import openpyxl
import pandas as pd
import pytest
from starlette.testclient import TestClient

from backend.app.core.ingestion import (
    inspect_dataset,
    parse_dataset_bytes,
    process_mapped_dataset,
)
from backend.app.main import app


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def sample_csv_bytes():
    csv_text = (
        "time,C_in,C_out,flow,volume\n"
        "0.0,1.0,0.1,1.0,10.0\n"
        "1.0,1.0,0.25,1.0,10.0\n"
        "2.0,1.0,0.35,1.0,10.0\n"
        "3.0,1.0,0.40,1.0,10.0\n"
        "4.0,1.0,0.40,1.0,10.0\n"
        "5.0,1.0,0.39,1.0,10.0\n"
        "6.0,1.0,0.41,1.0,10.0\n"
    )
    return csv_text.encode("utf-8")


@pytest.fixture
def sample_xlsx_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Run1"
    headers = ["time", "C_in", "C_out"]
    ws.append(headers)
    for t, cin, cout in [(0.0, 1.0, 0.1), (1.0, 1.0, 0.25), (2.0, 1.0, 0.4)]:
        ws.append([t, cin, cout])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Unit tests: parse_dataset_bytes & inspect_dataset
# ---------------------------------------------------------------------------


def test_parse_csv_bytes(sample_csv_bytes):
    df = parse_dataset_bytes(sample_csv_bytes, "experiment_run.csv")
    assert isinstance(df, pd.DataFrame)
    assert df.shape == (7, 5)
    assert list(df.columns) == ["time", "C_in", "C_out", "flow", "volume"]


def test_parse_excel_bytes(sample_xlsx_bytes):
    df = parse_dataset_bytes(sample_xlsx_bytes, "experiment_run.xlsx")
    assert isinstance(df, pd.DataFrame)
    assert df.shape == (3, 3)
    assert list(df.columns) == ["time", "C_in", "C_out"]


def test_parse_empty_bytes():
    with pytest.raises(ValueError, match="empty"):
        parse_dataset_bytes(b"", "empty.csv")


def test_parse_unsupported_extension():
    with pytest.raises(ValueError, match="Unsupported file format"):
        parse_dataset_bytes(b"data,here\n1,2", "malicious.exe")


def test_parse_max_size_exceeded():
    huge_data = b"a" * 1024
    with pytest.raises(ValueError, match="exceeds maximum allowed limit"):
        parse_dataset_bytes(huge_data, "large.csv", max_size=512)


def test_parse_max_columns_exceeded():
    cols = ",".join([f"col{i}" for i in range(10)])
    csv_bytes = f"{cols}\n{','.join(['1']*10)}".encode()
    with pytest.raises(ValueError, match="column count"):
        parse_dataset_bytes(csv_bytes, "wide.csv", max_cols=5)


def test_inspect_dataset(sample_csv_bytes):
    df = parse_dataset_bytes(sample_csv_bytes, "run1.csv")
    inspection = inspect_dataset(df, "run1.csv")
    assert inspection.filename == "run1.csv"
    assert inspection.file_type == "csv"
    assert inspection.total_rows == 7
    assert inspection.total_columns == 5
    assert len(inspection.preview_data) == 7
    assert inspection.duplicate_rows_count == 0


def test_inspect_dataset_detects_duplicates_and_missing():
    df = pd.DataFrame(
        {
            "t": [1.0, 1.0, 2.0],
            "val": [10.0, 10.0, None],
        }
    )
    inspection = inspect_dataset(df, "dup.csv")
    assert inspection.duplicate_rows_count == 1
    assert inspection.missing_counts["val"] == 1
    assert len(inspection.warnings) >= 2


# ---------------------------------------------------------------------------
# Unit tests: process_mapped_dataset
# ---------------------------------------------------------------------------


def test_process_mapped_dataset_conversion():
    df = pd.DataFrame(
        {
            "t_s": [0, 60, 120, 180, 240, 300],
            "C_out_raw": [0.1, 0.2, 0.4, 0.4, 0.4, 0.4],
            "C_in_raw": [1.0, 1.0, 1.0, 1.0, 1.0, 1.0],
        }
    )
    result = process_mapped_dataset(
        df=df,
        time_col="t_s",
        c_out_col="C_out_raw",
        c_in_col="C_in_raw",
        time_unit="s",
        conc_unit="mol/L",
    )
    # 0 s -> 0 min; 60 s -> 1 min
    assert result.series.time[1] == pytest.approx(1.0)
    # Conversion X_A = (1.0 - 0.4) / 1.0 = 0.60
    assert result.series.conversion[2] == pytest.approx(0.60)
    assert result.provenance_note.startswith("USER-PROVIDED DATA")
    assert not result.smoothing_applied
    assert not result.anomaly_detection_applied


def test_process_mapped_dataset_constant_inlet():
    df = pd.DataFrame(
        {
            "time": [0.0, 1.0, 2.0],
            "C_A": [0.4, 0.4, 0.4],
        }
    )
    result = process_mapped_dataset(
        df=df,
        time_col="time",
        c_out_col="C_A",
        constant_c_in=1.0,
        constant_flow=2.0,
        constant_volume=20.0,
    )
    # Conversion = (1.0 - 0.4)/1.0 = 0.60
    assert all(c == pytest.approx(0.60) for c in result.series.conversion)
    # Tau = V / Q = 20 / 2 = 10 min
    assert all(tau == pytest.approx(10.0) for tau in result.series.tau)


def test_conversion_undefined_when_cin_zero_or_negative():
    df = pd.DataFrame(
        {
            "time": [0.0, 1.0],
            "C_out": [0.2, 0.3],
            "C_in": [0.0, -1.0],
        }
    )
    result = process_mapped_dataset(
        df=df,
        time_col="time",
        c_out_col="C_out",
        c_in_col="C_in",
    )
    # When C_in <= 0, conversion must be None
    assert result.series.conversion[0] is None
    assert result.series.conversion[1] is None


def test_out_of_bounds_conversion_warning_not_deleted():
    # If C_out > C_in, conversion < 0. Data should NOT be deleted, only warned!
    df = pd.DataFrame(
        {
            "time": [0.0, 1.0],
            "C_out": [1.5, 0.5],
            "C_in": [1.0, 1.0],
        }
    )
    result = process_mapped_dataset(
        df=df,
        time_col="time",
        c_out_col="C_out",
        c_in_col="C_in",
    )
    assert len(result.series.time) == 2
    assert result.series.conversion[0] == pytest.approx(-0.5)
    assert any("conversion values lie outside physical range" in w for w in result.data_quality_warnings)


def test_non_monotonic_time_sorted():
    df = pd.DataFrame(
        {
            "time": [3.0, 1.0, 2.0],
            "C_out": [0.4, 0.2, 0.3],
            "C_in": [1.0, 1.0, 1.0],
        }
    )
    result = process_mapped_dataset(
        df=df,
        time_col="time",
        c_out_col="C_out",
        c_in_col="C_in",
    )
    assert result.series.time == [1.0, 2.0, 3.0]
    assert result.series.C_A_out_raw == [0.2, 0.3, 0.4]
    assert any("not monotonically increasing" in w for w in result.data_quality_warnings)


def test_savgol_smoothing_validation():
    df = pd.DataFrame(
        {
            "time": list(range(10)),
            "C_out": [0.5] * 10,
        }
    )
    # Even window length must raise
    with pytest.raises(ValueError, match="must be an odd integer"):
        process_mapped_dataset(
            df=df,
            time_col="time",
            c_out_col="C_out",
            apply_smoothing=True,
            savgol_window=4,
        )

    # Polyorder >= window length must raise
    with pytest.raises(ValueError, match="strictly less than window_length"):
        process_mapped_dataset(
            df=df,
            time_col="time",
            c_out_col="C_out",
            apply_smoothing=True,
            savgol_window=5,
            savgol_poly=5,
        )


def test_savgol_smoothing_execution():
    df = pd.DataFrame(
        {
            "time": list(range(15)),
            "C_out": [0.1, 0.12, 0.11, 0.13, 0.2, 0.25, 0.3, 0.35, 0.4, 0.41, 0.39, 0.4, 0.42, 0.4, 0.41],
        }
    )
    result = process_mapped_dataset(
        df=df,
        time_col="time",
        c_out_col="C_out",
        apply_smoothing=True,
        savgol_window=5,
        savgol_poly=2,
    )
    assert result.smoothing_applied is True
    assert result.series.C_A_out_smoothed is not None
    assert len(result.series.C_A_out_smoothed) == 15
    assert "caution" in result.smoothing_settings


def test_anomaly_detection_execution():
    # 20 regular points + 1 extreme spike
    y = [0.4] * 21
    y[10] = 5.0  # huge outlier
    df = pd.DataFrame(
        {
            "time": list(range(21)),
            "C_out": y,
        }
    )
    result = process_mapped_dataset(
        df=df,
        time_col="time",
        c_out_col="C_out",
        apply_anomaly_detection=True,
        contamination=0.05,
    )
    assert result.anomaly_detection_applied is True
    assert result.anomaly_count > 0
    assert result.series.is_anomaly is not None
    # Acceptance gate: points are NOT removed
    assert len(result.series.C_A_out_raw) == 21


# ---------------------------------------------------------------------------
# API Integration Tests
# ---------------------------------------------------------------------------


def test_api_upload_csv(client, sample_csv_bytes):
    response = client.post(
        "/api/v1/datasets/upload",
        files={"file": ("test.csv", sample_csv_bytes, "text/csv")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "test.csv"
    assert data["file_type"] == "csv"
    assert data["total_rows"] == 7
    assert "time" in data["columns"]
    assert len(data["raw_records"]) == 7


def test_api_upload_invalid_type(client):
    response = client.post(
        "/api/v1/datasets/upload",
        files={"file": ("malicious.bat", b"echo off", "text/plain")},
    )
    assert response.status_code == 422
    assert "Unsupported file format" in response.json()["detail"]


def test_api_process_dataset(client):
    records = [
        {"t": 0.0, "c_out": 0.1, "c_in": 1.0},
        {"t": 1.0, "c_out": 0.25, "c_in": 1.0},
        {"t": 2.0, "c_out": 0.40, "c_in": 1.0},
        {"t": 3.0, "c_out": 0.40, "c_in": 1.0},
        {"t": 4.0, "c_out": 0.40, "c_in": 1.0},
        {"t": 5.0, "c_out": 0.39, "c_in": 1.0},
        {"t": 6.0, "c_out": 0.41, "c_in": 1.0},
    ]
    payload = {
        "records": records,
        "time_col": "t",
        "c_out_col": "c_out",
        "c_in_col": "c_in",
        "time_unit": "min",
        "conc_unit": "mol/L",
        "apply_smoothing": True,
        "savgol_window": 5,
        "savgol_poly": 2,
        "apply_anomaly_detection": True,
        "contamination": 0.05,
    }
    response = client.post("/api/v1/datasets/process", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["smoothing_applied"] is True
    assert data["anomaly_detection_applied"] is True
    assert len(data["series"]["time"]) == 7
    assert data["provenance_note"].startswith("USER-PROVIDED DATA")


def test_api_sample_dataset(client):
    response = client.get("/api/v1/datasets/sample-dataset")
    assert response.status_code == 200
    data = response.json()
    assert data["total_rows"] == 101
    assert "time_min" in data["columns"]
    assert "C_A_outlet_mol_L" in data["columns"]
    assert len(data["raw_records"]) == 101


class TestCandidateDatasetsGovernanceAndSafety:
    """Verify Phase 6 requirements: Safe NPY loading, workbook profiling, and quarantine adherence."""

    def test_safe_npy_loading_with_allow_pickle_false(self):
        import numpy as np
        import os

        raw_dir = os.path.abspath("data/raw")
        npy_path = os.path.join(raw_dir, "cstr_rawdata.npy")
        if not os.path.exists(npy_path):
            pytest.skip("cstr_rawdata.npy not present in data/raw")

        # Must load with allow_pickle=False
        arr = np.load(npy_path, allow_pickle=False)
        assert arr.shape == (2860, 1404)
        assert arr.dtype == np.float64
        assert not np.isnan(arr).any()
        assert not np.isinf(arr).any()
        assert np.min(arr) > -50.0
        assert np.max(arr) < 1000.0

    def test_candidate_workbook_has_no_concentration(self):
        import os
        import openpyxl

        raw_dir = os.path.abspath("data/raw")
        xlsx_path = os.path.join(raw_dir, "CSTR_24aug_PRBS_dataset.xlsx")
        if not os.path.exists(xlsx_path):
            pytest.skip("CSTR_24aug_PRBS_dataset.xlsx not present in data/raw")

        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
        assert "CSTR Dataset" in wb.sheetnames
        ws = wb["CSTR Dataset"]

        first_rows = []
        for r_idx, row in enumerate(ws.iter_rows(max_row=5, values_only=True)):
            first_rows.append(row)
        wb.close()

        # Confirm no concentration labels (C_A, conc, mol/L) in headers
        header_text = " ".join(str(cell) for cell in first_rows[0] if cell is not None).lower()
        assert "c_a" not in header_text
        assert "concentration" not in header_text
        assert "mol/l" not in header_text
        # Thermal RTD sensor values are present
        assert "rtd" in header_text or "prbs" in header_text
